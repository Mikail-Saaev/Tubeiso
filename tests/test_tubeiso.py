"""Tests d'invariants. L'essentiel : un aller-retour LRA -> 3D -> developpe."""
import json
import math
import os
import re
import sys
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tubeiso import geometry
from tubeiso.model import Bend, TubeProgram
from tubeiso.parsers import crippa

PROG = """%MPF 412
(MASTERCUT 1.7/2.1/1.65, Ø6, SECTEUR, L=263, DS=69)
L56
L1 R1=155 R4=20 R6=263 R7=2 R12=39.5 R31=100 R33=280 M70
N1 L2 R15=46 R7=0
G91 G0 Y-35 G90 C0
G91 G0 Y-107 B24
N2 L3 R15=13 R41=11 R43=300
M30
"""


def test_parser():
    raw = crippa.parse(PROG)
    assert raw.name == "412"
    assert raw.tooling == "L56"
    assert raw.declared_length == 263.0
    assert raw.ds == 69.0
    assert raw.init["R12"] == 39.5
    assert raw.angles == [46.0, 13.0]
    assert raw.blocks[0].segment() == 142.0        # 35 + 107 [DOC 4.2]
    assert raw.head == 5 and raw.diameter == 6.0    # L56 = tete du bas, O6
    assert raw.blocks[0].rotation() == 24.0
    assert raw.complete


def test_truncation_detected():
    raw = crippa.parse("%MPF 1\n" + "x" * (255 - 7))
    assert any("255" in w for w in raw.warnings)
    assert not raw.complete


def test_developed_length_matches_geometry():
    """La polyligne 3D doit reproduire le developpe analytique."""
    tube = TubeProgram(
        ref="T", diameter=6.0, straights=[100.0, 80.0, 60.0],
        bends=[Bend(angle=45, rotation=0, clr=20),
               Bend(angle=90, rotation=90, clr=20)],
    )
    cl = geometry.build(tube, samples=400)
    analytic = geometry.developed_length(tube)
    assert abs(cl.developed - analytic) < 1e-9
    poly = sum(
        float(((cl.points[i + 1] - cl.points[i]) ** 2).sum() ** 0.5)
        for i in range(len(cl.points) - 1)
    )
    assert abs(poly - analytic) < 0.05          # erreur de discretisation


def test_planar_part_stays_planar():
    """Sans rotation, la piece reste dans un plan."""
    tube = TubeProgram(
        ref="T", diameter=6.0, straights=[50.0, 50.0, 50.0],
        bends=[Bend(angle=30, rotation=0, clr=15),
               Bend(angle=30, rotation=0, clr=15)],
    )
    cl = geometry.build(tube)
    assert abs(cl.points[:, 1]).max() < 1e-9


def test_missing_radius_refuses():
    tube = TubeProgram(ref="T", straights=[10.0, 10.0], bends=[Bend(angle=30)])
    try:
        geometry.build(tube)
    except geometry.MissingRadius:
        return
    raise AssertionError("aurait du lever MissingRadius")




def test_bsa_length_model_matches_excel():
    """Reproduit exactement Feuil2 de archivage_crippa.xlsm sur le tube 410."""
    from tubeiso import bsa
    straights = [4.0, 34.0, 63.5, 159.0]
    angles = [94.0, 34.0, 94.0]
    assert abs(bsa.arc_length(angles, 4) - 42.5993) < 1e-3
    r6 = bsa.developed_length(straights, angles, 4)
    assert abs(r6 - 301.82) < 0.01


def test_y_split_rule():
    """Y > 35 se scinde en 35 + reste. [XLSM Feuil1!E20]"""
    from tubeiso import bsa
    assert bsa.split_moves(34.0) == [34.0]
    assert bsa.split_moves(63.5) == [35.0, 28.5]
    assert bsa.split_moves(142.0) == [35.0, 107.0]


def test_tooling_code():
    from tubeiso import bsa
    assert bsa.parse_tooling("L54") == (5, 4)
    assert bsa.parse_tooling("L412") == (4, 12)



def test_primitives_are_exact():
    """Droites et arcs exacts : leur somme doit egaler le developpe."""
    tube = TubeProgram(
        ref="T", diameter=6.0, straights=[50.0, 40.0, 60.0],
        bends=[Bend(angle=45, rotation=0, clr=11),
               Bend(angle=90, rotation=90, clr=11)],
    )
    cl = geometry.build(tube)
    assert [p.kind for p in cl.primitives] == ["line", "arc", "line", "arc", "line"]
    assert abs(sum(p.length for p in cl.primitives) - cl.developed) < 1e-9


def test_solid_is_watertight():
    """Le solide balaye doit etre valide et creux. Ignore si cadquery absent."""
    from tubeiso import solid
    from tubeiso.model import Tooling
    try:
        solid._cq()
    except solid.SolidError:
        print("    (cadquery absent, test ignore)")
        return
    tube = TubeProgram(
        ref="T", diameter=6.0, straights=[50.0, 40.0, 60.0],
        bends=[Bend(angle=45, rotation=0, clr=11),
               Bend(angle=90, rotation=90, clr=11)],
    )
    cl = geometry.build(tube)
    rep = solid.report(tube, cl, Tooling(name="O6", diameter=6.0, clr=11, wall=1.0))
    assert rep["etanche"]
    # volume d'un tube creux : section annulaire fois developpe
    import math
    attendu = math.pi * (3.0 ** 2 - 2.0 ** 2) * cl.developed
    assert abs(rep["volume_mm3"] - attendu) / attendu < 0.01


# ---------------------------------------------------------------- v5 : lecture

def test_lft_multi_lignes_et_lots():
    """Un tube eclate sur plusieurs lignes est reconstitue, les lots separes."""
    from openpyxl import Workbook as XlWb
    from tubeiso import lft
    import tempfile, os

    wb = XlWb()
    ws = wb.active
    ws.append([]); ws.append(["export du jour"]); ws.append([])
    ws.append(["Repère", "CODE_MAT", "Long.", "LISTE", "PROGRAMME", "Prog Crippa",
               "EMBOUT_2", "RECOUPE_2"])
    ws.append([412, "293-421-006", 263, "LOT-A", "A-412", PROG, None, None])
    ws.append([412, None, None, "LOT-A", "A-412", None, "V06", None])   # 2e ligne
    ws.append([None, None, None, None, None, None, None, 7])            # continuation
    ws.append([412, "293-421-006", 263, "LOT-B", "B-412", PROG, "V06", None])
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    wb.save(path)

    book = lft.read(path)
    assert len(book.lots) == 2, "les deux lots doivent etre separes"
    a = book.lots[0].tubes[0]
    assert len(a.rows) == 3, "les 3 lignes du tube doivent etre regroupees"
    assert a.get("EMBOUT_2") == "V06", "aucune donnee ne doit etre perdue"
    assert a.recut == 7, "la recoupe de la ligne de continuation est conservee"
    assert a.list_number == "LOT-A" and a.program_number == "A-412"
    assert book.lots[1].tubes[0].list_number == "LOT-B"
    assert any("412" in w for w in book.warnings), "repere duplique a signaler"


def test_lft_entete_introuvable():
    from openpyxl import Workbook as XlWb
    from tubeiso import lft
    import tempfile, os
    wb = XlWb(); wb.active.append(["a", "b", "c"])
    path = os.path.join(tempfile.mkdtemp(), "x.xlsx")
    wb.save(path)
    try:
        lft.read(path)
    except ValueError:
        return
    raise AssertionError("un classeur sans colonnes LFT doit etre refuse")


# ------------------------------------------------------- v5 : retour elastique

def test_retour_elastique_aller_retour():
    """R15 -> angle reel -> R15 doit boucler, et donner des angles ronds.

    Hors de la fenetre d'equerre : dans cette fenetre plusieurs R15 decrivent
    le meme coude a 90 degres, donc l'aller-retour n'est plus bijectif et ne
    peut pas etre exige (voir test_equerre_a_90 juste apres).
    """
    from tubeiso import bsa
    cas = [(6, 46, 45), (6, 13, 13), (6, 39, 38), (6, 77, 75),
           (4, 34, 33), (4, 54, 52), (4, 26, 25),
           (8, 44, 43), (8, 31, 30), (8, 10, 10)]
    for d, r15, attendu in cas:
        assert bsa.locked_angle(r15) is None, f"cas mal choisi : R15={r15}"
        reel, delta = bsa.real_angle(r15, d)
        assert reel == attendu, f"Ø{d} R15={r15} -> {reel}, attendu {attendu}"
        assert bsa.programmed_angle(reel, d) == r15, f"Ø{d} : aller-retour casse"
        assert delta == r15 - attendu


def test_equerre_a_90():
    """Un R15 de 90 a 94 decrit une equerre, quel que soit le diametre.

    Le coefficient d'elasticite est une moyenne ; le programmeur ecrit ce qui
    sort de SA machine. Sur le corpus d'essai, R15=92 represente 39 % des
    coudes — ce sont des equerres, et les rendre a 89 ou 91 degres livrait une
    piece qui ne monte pas.
    """
    from tubeiso import bsa
    for d in (4, 6, 8, 10, 12, 15, 16, 18):
        for r15 in (90, 91, 92, 92.5, 93, 94):
            reel, delta = bsa.real_angle(r15, d)
            assert reel == 90.0, f"Ø{d} R15={r15} -> {reel}"
            assert delta == r15 - 90.0

    # hors fenetre, le calcul reprend la main
    assert bsa.real_angle(95, 6)[0] != 90.0
    assert bsa.real_angle(89, 6)[0] != 90.0
    assert bsa.locked_angle(94.0) == 90.0 and bsa.locked_angle(94.5) is None
    # le mode brut ne corrige rien, verrou compris
    assert bsa.real_angle(92, 8, "brut") == (92.0, 0.0)


def test_equerre_est_tracee_dans_la_piece():
    """Le verrou doit se voir : sur la piece, et dans les remarques."""
    from tubeiso import conventions, validate
    from tubeiso.model import Tooling

    prog = PROG.replace("R15=46", "R15=92").replace("R15=13", "R15=93")
    raw = crippa.parse(prog)
    tooling = Tooling("L56", 6.0, clr=11.0)
    tube = conventions.get("bsa").build(raw, tooling)
    assert [b.angle for b in tube.bends] == [90.0, 90.0]
    assert all(b.locked for b in tube.bends)
    assert any("verrouillé à 90" in w for w in tube.warnings)
    # information, jamais une anomalie : une equerre est une cote sure
    niveaux = {i.level for i in validate.check(tube, tooling)
               if "verrouillé" in i.message}
    assert niveaux == {validate.INFO}
    # et le classement ne doit pas dependre d'un accent
    assert validate._sans_accents("angle verrouillé à 90°") \
        .startswith("angle verrouille a 90")


def test_mode_brut_ne_corrige_rien():
    from tubeiso import bsa
    assert bsa.real_angle(94, 4, "brut") == (94.0, 0.0)


def test_geometrie_sur_angle_reel():
    """La 3D doit porter l'angle reel, le modele garde le R15 programme."""
    from tubeiso import bsa, conventions
    from tubeiso.model import Tooling
    raw = crippa.parse(PROG)
    tube = conventions.get("bsa").build(raw, Tooling("L56", 6.0, clr=11.0))
    assert [b.r15 for b in tube.bends] == [46.0, 13.0]
    assert [b.angle for b in tube.bends] == [45.0, 13.0]
    assert tube.bends[0].springback == 1.0


def test_tube_droit_a_une_longueur():
    """Un tube sans programme doit produire un cylindre, pas un point."""
    from tubeiso import conventions
    from tubeiso.model import Tooling
    tube = conventions.get("bsa").build_straight("500", 420.0, 6.0,
                                                 Tooling("Ø6", 6.0, clr=11.0))
    cl = geometry.build(tube)
    assert abs(cl.developed - 420.0) < 1e-9
    assert abs(float(np_norm(cl.points[-1] - cl.points[0])) - 420.0) < 1e-9


def test_controle_DS_detecte_un_faux_rayon():
    """Le controle de longueur doit reagir a un rayon errone. En v4 il valait
    exactement zero quoi qu'il arrive."""
    from tubeiso import conventions, validate
    from tubeiso.model import Tooling
    raw = crippa.parse(PROG)
    bon = conventions.get("bsa").build(raw, Tooling("L56", 6.0, clr=11.0))
    faux = conventions.get("bsa").build(raw, Tooling("L56", 6.0, clr=30.0))
    ecart = lambda t: abs(t.straights[-1] - t.ds)
    assert ecart(faux) > ecart(bon) + 5, "un rayon faux doit se voir"


def np_norm(v):
    return float((v ** 2).sum() ** 0.5)


# ---------------------------------------------------------------- v6 : matieres

def test_matiere_rigide_vs_souple():
    """La distinction rigide / souple vient du scan des codes matiere BSA."""
    from tubeiso import materials
    assert materials.kind_of("293-421-008") == materials.RIGIDE
    assert materials.kind_of("769-421-035") == materials.SOUPLE
    # meme prefixe 750, deux familles : le deuxieme triplet tranche
    assert materials.lookup("750 423 012").family == "Pneumatique"
    assert materials.lookup("750 421 012").family == "Forflex"
    assert materials.diameter("293-421-008") == 8
    assert materials.diameter("769-421-035") is None, "un flexible n'a pas de Ø cintrable"
    assert materials.lookup("293-421-008").wall == 1.0
    # famille deduite pour un code absent du catalogue
    inconnu = materials.lookup("778-421-099")
    assert inconnu is not None and inconnu.kind == materials.SOUPLE
    # Ø22 : rigide, mais BSA ne le cintre plus
    assert materials.lookup("293-421-022").kind == materials.RIGIDE
    assert not materials.lookup("293-421-022").bendable


def test_perimetre_progcrippa():
    """Seuls les tubes rigides porteurs d'une PROGCRIPPA sont traites."""
    from tubeiso import lft, scope
    from tubeiso.parsers import crippa

    class Rec:
        def __init__(self, code, iso="", droit=False, main=False, longueur=None):
            self.cells = {"CODE_MAT": code, "LONGUEUR": longueur}
            self.iso = iso
            self.straight = droit
            self.handmade = main

        def get(self, col, default=None):
            return self.cells.get(col, default)

        def number(self, col):
            v = self.cells.get(col)
            return float(v) if isinstance(v, (int, float)) else None

    # un tuyau souple SANS programme n'a pas de forme : fiche de debit, pas de
    # plan, pas de 3D — mais il n'est plus jete, sa matiere sert a commander.
    v = scope.evaluate(Rec("769-421-035", "", longueur=5000))
    assert v.ok and v.cut_only and not v.modelled
    assert v.reason == scope.MATIERE_SOUPLE and v.status == scope.DEBIT

    # meme code, mais AVEC un programme de cintrage : le programme fait foi,
    # et la contradiction est signalee en remarque.
    raw = crippa.parse(PROG)
    v = scope.evaluate(Rec("769-421-035", PROG), raw)
    assert v.ok and v.modelled and v.status == scope.TRAITE
    assert any("souple" in n for n in v.notes)

    # sans programme ET sans longueur, il n'y a rien a mettre sur le papier
    v = scope.evaluate(Rec("293-421-006", ""))
    assert not v.ok and v.reason == scope.LONGUEUR_ABSENTE

    # sans programme mais avec une longueur : tube droit, donc traite
    v = scope.evaluate(Rec("293-421-006", "", droit=True, longueur=300))
    assert v.ok and v.straight and v.reason == scope.TUBE_DROIT
    assert v.length == 300

    v = scope.evaluate(Rec("293-421-006", "", main=True, longueur=300))
    assert v.ok and v.cut_only and v.reason == scope.FAIT_MAIN

    v = scope.evaluate(Rec("293-421-006", PROG), raw)
    assert v.ok and v.diameter == 6 and v.kind == "rigide"

    # un programme tronque n'est plus rejete en bloc : il est traite, signale,
    # et c'est le controle de longueur qui decidera du modele 3D.
    tronque = PROG.replace("M30", "")
    v = scope.evaluate(Rec("293-421-006", tronque), crippa.parse(tronque))
    assert v.ok and v.modelled and v.reason == scope.PROGRAMME_TRONQUE


# --------------------------------------------------- v6 : correctifs de lecture

def test_cycle_de_chargement_L4():
    """[DOC 3.1] L4 charge les tubes longs : il porte R6 et R12 comme L1."""
    from tubeiso.parsers import crippa
    prog = PROG.replace("L1 R1=155", "L4 R1=155")
    raw = crippa.parse(prog)
    assert raw.loading == "L4"
    assert raw.init["R12"] == 39.5, "R12 doit etre lu sur un cycle L4"
    assert raw.declared_length == 263.0
    assert crippa.parse(PROG).loading == "L1"


def test_faux_pli_a_zero_degre_fusionne():
    """[DOC 5.3] Un R15=0 ne plie rien : les deux segments n'en font qu'un."""
    from tubeiso import bsa, conventions
    from tubeiso.model import Tooling
    from tubeiso.parsers import crippa

    # Repere 161 du lot Masterflex, tel quel : le dernier bloc est un L3 a
    # R15=0, et le DS du commentaire vaut le segment droit RECOLLE.
    prog = """%MPF 161
(Masterflex HD, Ø10, cassette, L=1444, ds=1306)
L510
L1 R1=155 R4=20 R6=1444 R7=2 R12=30.5 R31=100 R33=279.5 M70
N1 L2 R15=92
G91 G0 Y-35 G90 C0
G91 G0 Y-4 B90
N2 L2 R15=90
G91 G0 Y-35 G90 C0
G91 G0 Y-921
N3 L3 R15=0 R41=23 R43=300
M30
"""
    raw = crippa.parse(prog)
    tooling = Tooling(name="Ø10", diameter=10, clr=23, wall=1.0, elongation=5.0)
    tube = conventions.get("bsa").build(raw, tooling)
    assert tube.n_bends == 2, "le pli a 0° ne doit pas compter comme un coude"
    assert len(tube.straights) == 3
    assert tube.false_bends, "la fusion doit etre tracee"
    # le dernier segment recolle retombe sur le DS du commentaire, a l'arrondi
    assert abs(tube.straights[-1] - 1306) < 1.0, tube.straights


def test_r7_deja_present_ne_declenche_pas_d_alerte():
    """[XLSM Feuil2!B15] Le R7=0 vit dans un bloc N, pas dans la ligne L1."""
    from tubeiso import conventions, validate
    from tubeiso.model import Tooling
    from tubeiso.parsers import crippa

    raw = crippa.parse(PROG)            # le 412 porte R7=0 dans le bloc N1
    tooling = Tooling(name="Ø6", diameter=6, clr=11, wall=1.0, elongation=5.0,
                      min_straight=11)
    tube = conventions.get("bsa").build(raw, tooling)
    assert tube.r7_released
    codes = [i.code for i in validate.check(tube, tooling)]
    assert "r7_manquant" not in codes


def test_angles_au_demi_degre():
    """Le programmeur ecrit au demi-degre : 46.5 en Ø6 doit rester lisible."""
    from tubeiso import bsa
    assert bsa.real_angle(46.5, 6)[0] == 45.5
    assert bsa.real_angle(46, 4)[0] == 45.0, "l'exemple meme de [DOC 5.4]"
    assert bsa.real_angle(25.5, 6)[0] == 25.0
    assert bsa.real_angle(92.5, 8)[0] == 90.0, "un R15 de 92.5 reste une équerre"
    # L'aller-retour reste exact sur tous les angles entiers, SAUF ceux que le
    # verrou d'equerre ramene volontairement a 90 : la, plusieurs R15 mènent au
    # meme angle, et c'est le but.
    for d in (4, 6, 8, 10, 12, 15, 16, 18):
        for theta in range(5, 186):
            r15 = bsa.programmed_angle(theta, d)
            if bsa.locked_angle(r15) is not None:
                assert bsa.real_angle(r15, d)[0] == 90.0
                assert abs(theta - 90) <= 3, (d, theta, r15)
                continue
            assert bsa.real_angle(r15, d)[0] == theta, (d, theta, r15)


def test_tables_du_classeur_pour_le_diametre_16():
    """[XLSM Feuil2!B47] Ø16 : droite mini 36 mm, et non 30 comme le Ø15."""
    from tubeiso import bsa
    assert bsa.MIN_STRAIGHT[16] == 36
    assert bsa.R15_FOR_90[16] == 92.5, "sans quoi le Ø16 sort sans compensation"
    assert bsa.real_angle(92.5, 16)[0] == 90.0


# ------------------------------------------------------------ v6 : rattachement

def test_rattachement_groupe_machine():
    from tubeiso import registry
    e = registry.parse_filename("BCH_PLATINE_82_0889_0877-0000-CL")
    assert e.groupe == "BSH", "le prefixe de fichier BCH designe le groupe BSH"
    assert e.machine == "PLATINE_82_0889"
    assert registry.list_number_from("BCH_PLATINE_82_0889_0877-0000-CL") == "0877-0000-CL"
    assert registry.safe_name("A/B:C*?") == "A_B_C"
    assert registry.parse_filename("FLX_MASTERFLEX_HD_EXPERFLEX_0227_0198-0002-JV").groupe == "FLX"


# ------------------------------------------------------------------ v6 : le plan

def test_plan_pdf_autoportant():
    """Le PDF doit exister, faire deux pages, et porter les donnees cles."""
    import tempfile

    from tubeiso import conventions, geometry, render, validate
    from tubeiso.model import Tooling
    from tubeiso.parsers import crippa

    raw = crippa.parse(PROG)
    tooling = Tooling(name="Ø6", diameter=6, clr=11, wall=1.0, elongation=5.0,
                      material="Tube Ermeto zingué 6/4", min_straight=11)
    tube = conventions.get("bsa").build(raw, tooling)
    cl = geometry.build(tube)
    issues = validate.check(tube, tooling, cl)
    data = render.PlanData(tube=tube, centerline=cl, tooling=tooling, issues=issues,
                           groupe="FLX", machine="ESSAI", lft="ESSAI_0000-0000-XX")

    svg = render.to_svg(data)
    assert svg.startswith("<svg") and "PLAN DE FABRICATION" in svg
    assert "TABLE LRA" not in svg, "la page 1 ne porte pas la table LRA"

    with tempfile.TemporaryDirectory() as tmp:
        out = render.to_pdf(data, Path(tmp) / "412.pdf")
        blob = out.read_bytes()
        assert blob.startswith(b"%PDF"), "en-tete PDF absent"
        assert blob.count(b"/Type /Page") >= 2 or blob.count(b"/Type/Page") >= 2, \
            "le plan doit faire deux pages"
        assert len(blob) > 3000

        cahier = render.booklet([data], Path(tmp) / "cahier.pdf", "ESSAI",
                                [{"repere": "9", "motif": "matière_souple",
                                  "detail": "Uniflex 44/35"}])
        assert cahier.read_bytes().startswith(b"%PDF")


def test_segments_du_plan_suivent_le_tube():
    """Les cotes du plan doivent porter sur les vrais segments droits."""
    from tubeiso import geometry, render
    from tubeiso.model import Bend, TubeProgram

    tube = TubeProgram(ref="T", diameter=8.0, straights=[50.0, 60.0, 70.0],
                       bends=[Bend(angle=90, rotation=0, clr=14),
                              Bend(angle=45, rotation=90, clr=14)])
    cl = geometry.build(tube)
    spans = render.straight_spans(cl, tube.n_bends)
    assert len(spans) == len(tube.straights)
    for (a, b), attendu in zip(spans, tube.straights):
        longueur = float(((b - a) ** 2).sum() ** 0.5)
        assert abs(longueur - attendu) < 1e-6, (longueur, attendu)


# ------------------------------------------------------------- v6 : campagne

def test_campagne_range_et_indexe():
    """Une campagne produit l'arborescence, l'index et le rapport."""
    import tempfile

    from openpyxl import Workbook, load_workbook

    from tubeiso import batch

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sources"
        src.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.append(["REP", "CODE_MAT", "LONGUEUR", "LISTE", "PROGRAMME",
                   "PROGCRIPPA", "DROIT"])
        ws.append(["412", "293-421-006", 263, "0792-0002-JV", "792_JV-412",
                   PROG, "Faux"])
        ws.append(["900", "769-421-035", 5000, "0792-0002-JV", "", "", "Vrai"])
        ws.append(["901", "293-421-006", 300, "0792-0002-JV", "", "", "Vrai"])
        wb.save(src / "BCH_ESSAI_0001_0792-0002-JV.xlsx")

        out = Path(tmp) / "biblio"
        campagne = batch.run([src], out, batch.Options(models=False))
        s = campagne.summary()
        # 412 est cintre, 901 est un tube droit rigide, 900 est un flexible :
        # il n'a pas de forme, mais il a une matiere et une longueur, donc une
        # fiche de debit — et surtout pas un plan.
        assert s["pieces"] == 3
        assert s["traitees"] == 1 and s["tubes_droits"] == 1
        assert s["debits"] == 1 and s["exclues"] == 0
        assert s["plans"] == 2 and s["fiches_debit"] == 1

        # Arborescence a plat : un dossier par type, pas de niveaux imbriques.
        base = "BCH_ESSAI_0001_0792-0002-JV"
        assert (out / "plans" / f"{base}_412.pdf").exists()
        assert (out / "plans" / f"{base}_901.pdf").exists(), "le tube droit a un plan"
        assert (out / "donnees" / f"{base}_412.json").exists()
        assert (out / "cahiers" / f"{base}_cahier.pdf").exists()
        assert not (out / "plans" / f"{base}_900.pdf").exists(), \
            "un tuyau souple ne doit produire aucun plan"
        fiche = out / "debits" / f"{base}_900.pdf"
        assert fiche.exists(), "le tuyau souple sort en fiche de debit"
        assert b"FICHE DE D" in fiche.read_bytes()[:200] or fiche.stat().st_size > 800
        dossiers = {d.name for d in out.iterdir() if d.is_dir()}
        assert dossiers == {"plans", "debits", "donnees", "cahiers"}, dossiers
        assert not any(d.is_dir() for d in (out / "plans").iterdir()), \
            "aucun sous-dossier sous plans/"

        donnees = json.loads(
            (out / "donnees" / f"{base}_412.json").read_text("utf-8"))
        assert donnees["matiere"]["nature"] == "rigide"
        assert donnees["cintrage"]["lra"][0]["r15_programme"] == 46
        assert len(donnees["geometrie"]["sommets_xyz"]) == tube_sommets(donnees)

        souple = json.loads(
            (out / "donnees" / f"{base}_900.json").read_text("utf-8"))
        assert souple["forme"]["definie"] is False
        assert souple["debit"]["longueur_a_debiter"] == 5000
        assert "cintrage" not in souple and "geometrie" not in souple

        index = load_workbook(out / "INDEX.xlsx")
        assert index.sheetnames == ["Tubes", "LFT", "Campagne"]
        assert index["Tubes"].max_row == 4          # en-tete + 3 pieces
        motifs = {index["Tubes"].cell(row=r, column=9).value
                  for r in range(2, 5)}
        assert "matière_souple" in motifs
        statuts = {index["Tubes"].cell(row=r, column=8).value for r in range(2, 5)}
        assert statuts == {"traitée", "tube droit", "débit seul"}

        # reprise : la meme campagne relancee ne refait rien
        assert (out / batch.STATE_FILE).exists(), "l'etat de reprise doit exister"
        again = batch.run([src], out, batch.Options(models=False))
        assert again.summary()["fichiers_sautes"] == 1
        # ... sauf si on le demande
        forcee = batch.run([src], out, batch.Options(models=False, force=True))
        assert forcee.summary()["fichiers_sautes"] == 0


def tube_sommets(donnees: dict) -> int:
    return len(donnees["cintrage"]["lra"]) + 2


# ------------------------------------------------- v6.1 : PDF sans dependance

def test_pdf_ecrit_sans_bibliotheque():
    """Le PDF ne doit dependre que de la bibliotheque standard.

    La premiere version passait par reportlab, et l'export tombait en panne sur
    tout poste ou il n'etait pas installe. Ce test verrouille la sortie : entete,
    pagination, table des objets, et metriques de police exactes.
    """
    import tempfile

    from tubeiso import sheet

    source = Path(sheet.__file__).read_text(encoding="utf-8")
    assert "import reportlab" not in source and "from reportlab" not in source

    # metriques Adobe officielles : un texte cale a droite en depend
    assert abs(sheet.text_width("Ø8 × 1 mm", "Helvetica", 10) - 49.74) < 0.01
    assert abs(sheet.text_width("ABC", "Helvetica-Bold", 12) - 25.992) < 0.01

    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "essai.pdf"
        s = sheet.PdfSheet(out, title="Essai — accentué")
        s.rect(10, 10, 100, 50)
        s.text(20, 30, "Ø18 × 2 mm — cintrage à 92,5°", 4.0)
        s.polyline([(10, 60), (50, 80), (90, 60)], w=0.5)
        s.circle(60, 40, 5, fill="#EEEEEE")
        s.page_break()
        s.text(20, 20, "page 2", 4.0)
        s.save()
        blob = out.read_bytes()

    assert blob.startswith(b"%PDF-1.4")
    assert blob.rstrip().endswith(b"%%EOF")
    assert blob.count(b"/Type /Page ") == 2, "deux pages attendues"
    assert b"/Type /Pages" in blob and b"/Type /Catalog" in blob
    assert b"xref" in blob and b"startxref" in blob
    assert b"/FlateDecode" in blob, "les flux doivent etre compresses"

    # la table xref doit pointer sur de vrais objets
    tail = blob[blob.rindex(b"startxref"):]
    start = int(tail.split(b"\n")[1])
    assert blob[start:start + 4] == b"xref"


def test_largeur_de_texte_identique_svg_et_pdf():
    """Les deux sorties doivent tronquer au meme endroit, sinon l'apercu ment."""
    from tubeiso import sheet
    svg = sheet.SvgSheet()
    for txt, bold in (("Ø8 × 1 mm", False), ("MASTERFLEX_HD_EXPERFLEX", True),
                      ("92,5°", False)):
        a = svg.width_of(txt, 2.5, bold)
        b = sheet.text_width(txt, "Helvetica-Bold" if bold else "Helvetica",
                             2.5 * sheet.PT_PER_MM) * sheet.MM_PER_PT
        assert abs(a - b) < 1e-9


# --------------------------------------------- v6.1 : tubes droits exportables

def test_tube_rigide_sans_programme_est_traite():
    """[Info Crippa] « meme s'il n'y a pas de programme, il faut generer la 3D
    avec uniquement la longueur et le diametre »."""
    from tubeiso import materials, scope

    class Rec:
        def __init__(self, code, longueur, droit=False, main=False, iso=""):
            self.cells = {"CODE_MAT": code, "LONGUEUR": longueur}
            self.iso, self.straight, self.handmade = iso, droit, main

        def get(self, col, default=None):
            return self.cells.get(col, default)

        def number(self, col):
            v = self.cells.get(col)
            return float(v) if isinstance(v, (int, float)) else None

    v = scope.evaluate(Rec("293-421-008", 640, droit=True))
    assert v.ok and v.straight and v.status == scope.DROIT and v.diameter == 8

    # sans longueur, il n'y a rien a modeliser
    v = scope.evaluate(Rec("293-421-008", None, droit=True))
    assert not v.ok and v.reason == scope.LONGUEUR_ABSENTE

    # un souple n'a pas de forme, meme droit : fiche de debit, pas de modele
    v = scope.evaluate(Rec("769-421-035", 5000, droit=True))
    assert v.ok and v.cut_only and v.reason == scope.MATIERE_SOUPLE

    # un rigide hors outillage de cintrage se coupe droit : BSA ne le plie pas,
    # donc l'absence de matrice ne l'empeche pas d'etre modelise.
    v = scope.evaluate(Rec("293-421-028", 900, droit=True))
    assert v.ok and v.straight and v.diameter == 28 and v.wall == 3.0


def test_plan_d_un_tube_droit():
    import tempfile

    from tubeiso import conventions, geometry, render, validate
    from tubeiso.model import Tooling

    tooling = Tooling(name="Ø15", diameter=15, clr=45, wall=1.5, elongation=4.0)
    tube = conventions.get("bsa").build_straight("100", 908.0, 15, tooling)
    cl = geometry.build(tube)
    assert abs(cl.developed - 908.0) < 1e-9
    data = render.PlanData(tube=tube, centerline=cl, tooling=tooling,
                           issues=validate.check(tube, tooling, cl),
                           lft="FLX_ESSAI_0000-0000-XX")
    svg = render.to_svg(data)
    assert "TUBE DROIT" in svg and "sans objet" in svg
    with tempfile.TemporaryDirectory() as tmp:
        assert render.to_pdf(data, Path(tmp) / "d.pdf").read_bytes().startswith(b"%PDF")


# ------------------------------------------------ v6.1 : nommage des exports

def test_nom_de_fichier_tracable():
    from tubeiso import registry
    assert registry.output_basename(
        "BCH_PLATINE_82_0889_0877-0000-CL.xlsx", "170"
    ) == "BCH_PLATINE_82_0889_0877-0000-CL_170"
    assert registry.output_basename("BCH_X_0001-0000-AA", "105.1") == \
        "BCH_X_0001-0000-AA_105.1"
    assert registry.output_basename("", "42") == "42"
    assert registry.output_basename("LFT", "a/b:c") == "LFT_a_b_c"


# --------------------------------------------- v6.2 : correctifs de l'audit

def test_code_matiere_rendu_comme_un_nombre():
    """openpyxl rend un code saisi sans tiret comme un nombre, parfois flottant."""
    from tubeiso import materials
    from tubeiso.config import code_mat_diameter
    for valeur in ("293-421-008", "293 421 008", "BSA293421008",
                   293421008, 293421008.0):
        assert materials.digits(valeur) == "293421008", valeur
        assert materials.diameter(valeur) == 8, valeur
        assert code_mat_diameter(valeur) == 8, valeur


def test_recherche_d_azimut_ne_depend_pas_du_nombre_de_points():
    """Comparer toutes les paires de points coutait une seconde par piece."""
    import time

    from tubeiso import geometry, render
    from tubeiso.model import Bend, TubeProgram

    temps = []
    for n in (5, 40):
        tube = TubeProgram(ref="T", diameter=6.0, straights=[25.0] * (n + 1),
                           bends=[Bend(angle=45, rotation=90, clr=11)
                                  for _ in range(n)])
        cl = geometry.build(tube)
        t0 = time.perf_counter()
        render.best_azimuth(cl, 6.0)
        geometry.min_segment_distance(cl, 6.0, 11.0)
        temps.append(time.perf_counter() - t0)
    assert temps[1] < temps[0] * 4, (
        f"le coût explose avec le nombre de coudes : {temps}")


def test_tube_droit_sans_diametre_est_refuse():
    """Un plan qui annonce Ø0 n'est pas fabricable : il doit lever une erreur."""
    from tubeiso import conventions, validate
    from tubeiso.model import Tooling

    tl = Tooling(name="?", diameter=0.0)
    tube = conventions.get("bsa").build_straight("X", 500.0, 0.0, tl)
    codes = {i.code: i.level for i in validate.check(tube, tl)}
    assert codes.get("diametre_absent") == validate.ERROR

    tl8 = Tooling(name="Ø8", diameter=8, clr=14, wall=1.0)
    court = conventions.get("bsa").build_straight("Z", 12.0, 8.0, tl8)
    assert "tube_droit_court" in {i.code for i in validate.check(court, tl8)}


def test_piece_sans_forme_sort_une_fiche_de_debit():
    """Le serveur ne doit jamais inventer un modèle pour une pièce sans forme.

    Un tuyau souple n'a ni rayon, ni angles, ni géométrie : il sort une fiche
    de débit, dans son propre dossier, et l'export 3D est refusé avec un motif.
    """
    import tempfile

    from openpyxl import Workbook

    from tubeiso.app.server import Session, create_app

    with tempfile.TemporaryDirectory() as tmp:
        wb = Workbook()
        ws = wb.active
        ws.append(["REP", "CODE_MAT", "LONGUEUR", "LISTE", "PROGRAMME",
                   "PROGCRIPPA", "DROIT"])
        ws.append(["9", "769-421-050", 5000, "L", "P", "", "Vrai"])
        ws.append(["412", "293-421-006", 263, "L", "P2", PROG, "Faux"])
        source = Path(tmp) / "BCH_X_0001-0000-AA.xlsx"
        wb.save(source)

        s = Session()
        s.open_lft(str(source))
        c = create_app(s).test_client()
        tubes = [t for l in c.get("/api/tubes").get_json()["lots"]
                 for t in l["tubes"]]
        souple = next(t for t in tubes if t["ref"] == "9")
        assert souple["scope"] == "débit seul" and souple["nature"] == "souple"
        assert souple["cut_only"] is True and souple["status"] == "débit"
        assert souple["livrable"] == "fiche de débit"

        d = c.get(f"/api/tube/{souple['uid']}").get_json()
        assert d["out_of_scope"] is True and d["cut_only"] is True
        assert d["polyline"] == [] and d["bends"] == [] and d["mesh"] is None
        assert "developed" not in d, "aucun développé pour une pièce sans forme"
        assert d["declared"] == 5000, "la longueur reste disponible pour le débit"

        # l'export ecrit la fiche, et refuse le 3D avec un motif
        r = c.post("/api/export", json={"uids": [souple["uid"]], "dir": tmp,
                                        "formats": ["step", "pdf"]}).get_json()
        assert len(r["written"]) == 1 and r["written"][0].endswith(".pdf")
        assert "debits" in r["written"][0], r["written"][0]
        assert r["failed"] and "forme non définie" in r["failed"][0]["error"]
        assert "STEP" in r["failed"][0]["error"]


def test_export_range_par_type():
    """L'export manuel suit la même arborescence plate que la campagne."""
    import tempfile

    from tubeiso import batch

    assert batch.folder_for("pdf") == "plans"
    assert batch.folder_for("debit") == "debits"
    assert batch.folder_for(".stp") == "step"
    assert batch.folder_for("json") == "donnees"
    assert batch.folder_for("zzz") == "autres"


# ------------------------------- v6.3 : le modèle 3D suit la fiabilité du plan

def test_programme_tronque_sans_perte_garde_son_modele():
    """Une troncature qui ne mange que la fin de ligne ne coûte rien.

    L'ancienne règle rejetait tout programme sans M30, alors qu'un `M30`
    manquant ne retire aucune géométrie : le développé recalculé referme
    l'équation de longueur au millimètre près.
    """
    from tubeiso import conventions, geometry, validate
    from tubeiso.config import Config
    from tubeiso.parsers import crippa

    cfg = Config.load(None)
    raw = crippa.parse(PROG.replace("M30", ""))
    assert not raw.complete, "le programme doit bien être vu comme tronqué"
    tooling = cfg.for_program(raw.tooling, raw.diameter, "293-421-006")
    tube = conventions.get("bsa").build(raw, tooling, angle_mode=cfg.angle_mode)
    cl = geometry.build(tube)
    issues = validate.check(tube, tooling, centerline=cl)
    codes = {i.code: i.level for i in issues}
    assert codes.get("programme_tronque") == validate.WARN
    assert "programme_tronque_perte" not in codes
    assert validate.unsafe(issues) == "", "le modèle 3D reste autorisé"


def test_programme_tronque_avec_perte_perd_son_modele():
    """Une troncature qui mange des blocs interdit le STEP, pas le plan.

    Un solide faux part chez un sous-traitant sans que personne ne relise le
    plan : c'est le seul cas où l'application refuse d'écrire un modèle.
    """
    from tubeiso import conventions, geometry, validate
    from tubeiso.config import Config
    from tubeiso.parsers import crippa

    cfg = Config.load(None)
    entier = crippa.parse(PROG)
    # on coupe le dernier bloc de cintrage : le tube perd un coude et sa droite
    lignes = [l for l in PROG.splitlines() if l.strip()]
    coupe = "\n".join(lignes[:-2])
    raw = crippa.parse(coupe)
    assert len(raw.blocks) < len(entier.blocks), "il faut vraiment perdre un bloc"

    tooling = cfg.for_program(raw.tooling, raw.diameter, "293-421-006")
    tube = conventions.get("bsa").build(raw, tooling, angle_mode=cfg.angle_mode)
    tube.declared_length = entier.declared_length      # R6 reste celui du tout
    cl = geometry.build(tube)
    issues = validate.check(tube, tooling, centerline=cl)
    assert validate.unsafe(issues) == "programme_tronque_perte"
    perte = next(i for i in issues if i.code == "programme_tronque_perte")
    assert perte.level == validate.ERROR
    assert "mm" in perte.message, "le message doit chiffrer ce qui manque"


def test_campagne_refuse_le_step_mais_ecrit_le_plan():
    """Bout en bout : plan écrit, bandeau ERREUR, aucun STEP."""
    import tempfile

    from openpyxl import Workbook

    from tubeiso import batch

    lignes = [l for l in PROG.splitlines() if l.strip()]
    coupe = "\n".join(lignes[:-2])

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sources"
        src.mkdir()
        wb = Workbook()
        ws = wb.active
        ws.append(["REP", "CODE_MAT", "LONGUEUR", "LISTE", "PROGRAMME",
                   "PROGCRIPPA", "DROIT"])
        ws.append(["7", "293-421-006", 263, "0001-0000-AA", "P", coupe, "Faux"])
        wb.save(src / "BCH_ESSAI_0002_0001-0000-AA.xlsx")

        out = Path(tmp) / "biblio"
        campagne = batch.run([src], out, batch.Options())
        row = campagne.rows[0]
        assert row.statut == "traitée" and row.motif == "programme_tronqué"
        assert row.plan_pdf, "le plan doit être écrit"
        assert not row.modele_3d, "le STEP doit être refusé"
        assert row.sans_3d == "programme_tronque_perte"
        assert row.controle == "ERREUR"
        assert campagne.summary()["sans_3d"] == 1
        assert not (out / "step").exists()


def test_couverture_du_cahier_est_paginee():
    """Un lot de soixante pièces ne tient pas sur une page.

    L'ancienne couverture écrivait toutes les lignes à la suite : les
    dernières sortaient du cadre et le bloc « pièces sans plan » n'était
    jamais imprimé. Il est en bas de la couverture, donc c'est exactement
    celui qu'on perdait.
    """
    import tempfile

    from tubeiso import conventions, geometry, render, validate
    from tubeiso.model import Tooling

    tooling = Tooling(name="Ø8", diameter=8, clr=14, wall=1.0, elongation=4.0)
    conv = conventions.get("bsa")
    items = []
    for i in range(60):
        tube = conv.build_straight(str(100 + i), 500.0 + i, 8, tooling)
        cl = geometry.build(tube)
        items.append(render.PlanData(tube=tube, centerline=cl, tooling=tooling,
                                     issues=validate.check(tube, tooling, cl),
                                     lft="BCH_ESSAI_0000-0000-XX"))
    exclus = [{"repere": "9", "motif": "matière_souple",
               "detail": "tuyau souple", "fiche": True}]

    with tempfile.TemporaryDirectory() as tmp:
        pdf = render.booklet(items, Path(tmp) / "c.pdf", lot_label="ESSAI",
                             excluded=exclus)
        data = pdf.read_bytes()
        assert data.startswith(b"%PDF")
        pages = data.count(b"/Type /Page\n") or data.count(b"/Type /Page")
        # 60 pièces × 2 pages + au moins 2 pages de couverture
        assert pages >= 122, pages

    # la couverture seule, en SVG, pour lire ce qu'elle contient vraiment
    s = render.sh.SvgSheet(render.PAGE_W, render.PAGE_H)
    render._cover(s, items, "ESSAI", exclus)
    textes = [s.to_svg(i) for i in range(3)]
    assert "CAHIER DE FABRICATION" in textes[0]
    assert "(suite)" in textes[1], "la seconde page doit porter la mention suite"
    joint = " ".join(textes)
    assert "SANS PLAN DE CINTRAGE" in joint
    assert "fiche de d" in joint, "la sortie de la pièce doit être nommée"
    # aucune ligne ne doit être écrite sous le cadre
    for page in textes:
        for y in re.findall(r'y="([0-9.]+)"', page):
            assert float(y) <= render.FRAME[3] + 0.5, y


def test_le_plan_dit_lui_meme_qu_il_n_a_pas_de_modele():
    """Le plan part seul chez le sous-traitant : c'est lui qui doit le dire."""
    from tubeiso import conventions, geometry, render, validate
    from tubeiso.config import Config
    from tubeiso.parsers import crippa

    cfg = Config.load(None)
    entier = crippa.parse(PROG)
    lignes = [l for l in PROG.splitlines() if l.strip()]
    raw = crippa.parse("\n".join(lignes[:-2]))
    tooling = cfg.for_program(raw.tooling, raw.diameter, "293-421-006")
    tube = conventions.get("bsa").build(raw, tooling, angle_mode=cfg.angle_mode)
    tube.declared_length = entier.declared_length
    cl = geometry.build(tube)
    data = render.PlanData(tube=tube, centerline=cl, tooling=tooling,
                           issues=validate.check(tube, tooling, centerline=cl),
                           lft="BCH_ESSAI_0000-0000-XX")
    assert data.status == "ERREUR"
    svg = render.to_svg(data)
    assert "AUCUN MODÈLE 3D N" in svg
    # le chiffre de la matière manquante ne doit pas être coupé
    assert "mm de matière manquante" in svg
    assert "…" not in svg.split("AUCUN MODÈLE")[0].split("CONTRÔLE EN ERREUR")[-1]


# ------------------------------------- v6.3 : sens de l'axe B (B+ = horaire)

def test_rotation_B_positive_est_horaire():
    """B+90 tourne dans le sens des aiguilles d'une montre, B-90 dans l'autre.

    L'observateur est a l'extremite aval et regarde le tube revenir vers la
    machine. Jusqu'a la v6.3 le code appliquait l'inverse : le cartouche du
    plan annonçait « horaire » et la piece sortait en miroir, a l'ecran comme
    au format STEP.

    Le test raisonne sur les vecteurs, pas sur une image : avec un seul coude
    de 90 degres, le tube part suivant +X et se couche vers +Z quand B=0. Pour
    l'observateur place en aval (+X pointe vers lui), +Z est « en haut » et +Y
    est « a droite ». Passer de haut a droite, c'est tourner dans le sens
    horaire — donc B=+90 doit amener l'extremite suivant +Y.
    """
    import numpy as np

    from tubeiso import geometry
    from tubeiso.model import Bend, TubeProgram

    def sortie(rotation, hand=1):
        t = TubeProgram(ref="B", diameter=8, straights=[50.0, 50.0],
                        bends=[Bend(angle=90.0, rotation=rotation, clr=14.0)])
        cl = geometry.build(t, handedness=hand)
        v = cl.points[-1] - cl.points[-2]
        return v / np_norm(v)

    haut, droite = np.array([0, 0, 1.0]), np.array([0, 1.0, 0])
    assert np_norm(sortie(0) - haut) < 1e-9, "sans rotation, le tube se couche vers +Z"
    assert np_norm(sortie(+90) - droite) < 1e-9, "B+90 : de haut vers la droite = horaire"
    assert np_norm(sortie(-90) + droite) < 1e-9, "B-90 : vers la gauche = antihoraire"

    # le reglage d'atelier refait bien la piece miroir
    assert np_norm(sortie(+90, hand=-1) + droite) < 1e-9

    # un demi-tour est un demi-tour, quel que soit le signe [DOC 8.3.5]
    assert np_norm(sortie(+180) - sortie(-180)) < 1e-9


def test_step_relu_rend_les_memes_rotations():
    """Aller-retour complet : LRA -> solide -> STEP -> LRA.

    C'est le seul controle qui prouve que le fichier livre au sous-traitant
    tourne du meme cote que le programme d'origine. Le lecteur de STEP mesure
    un angle direct autour de la tangente ; sans le facteur geometry.B_SIGN il
    rendait des rotations de signe oppose.
    """
    import tempfile

    from tubeiso import geometry, solid, stepreader
    from tubeiso.model import Bend, Tooling, TubeProgram

    rotations = [0.0, 90.0, -90.0, 45.0, -135.0]
    tube = TubeProgram(ref="RT", diameter=8.0,
                       straights=[60.0, 55.0, 55.0, 55.0, 55.0, 60.0],
                       bends=[Bend(angle=90.0, rotation=r, clr=14.0)
                              for r in rotations])
    tooling = Tooling("Ø8", 8.0, clr=14.0, wall=1.0)
    cl = geometry.build(tube)
    with tempfile.TemporaryDirectory() as tmp:
        f = solid.export(tube, cl, Path(tmp), tooling, ["step"], basename="rt")[0]
        lra = stepreader.analyse(str(f))["lra"]
    assert [round(v, 2) for v in lra["rotations"]] == rotations
    assert [round(v, 2) for v in lra["angles"]] == [90.0] * 5
    assert [round(v, 2) for v in lra["segments"]] == tube.straights


def _run() -> int:
    echecs = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_"):
            continue
        try:
            fn()
        except Exception as exc:                       # noqa: BLE001
            echecs += 1
            print(f"ECHEC  {name} : {type(exc).__name__}: {exc}")
            traceback.print_exc()
        else:
            print(f"ok  {name}")
    print("\nTous les tests passent." if not echecs
          else f"\n{echecs} test(s) en echec.")
    return 1 if echecs else 0


if __name__ == "__main__":
    code = _run()
    # Sortie immediate, sans passer par la fermeture de l'interpreteur.
    # OpenCascade et VTK liberent leurs ressources natives a ce moment-la et
    # y plantent parfois, ce qui renvoie un code non nul malgre des tests
    # tous verts. Le symptome a ete observe sur les runners GitHub.
    sys.stdout.flush()
    sys.stderr.flush()
    os._exit(code)


