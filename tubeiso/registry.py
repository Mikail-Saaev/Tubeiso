"""Rattachement d'une LFT a son groupe et a sa machine.

Le classement demande est : **Groupe → Machine → LFT → plan → 3D → données.**
Encore faut-il savoir a quelle machine appartient un fichier LFT. Deux sources,
dans cet ordre :

1. `Repertoire_Machines_Consolide.xlsm`, qui fait autorite. Sa feuille
   `Repertoire_LFT` associe chaque code LFT a son groupe, sa machine et sa
   description. Le code LFT est exactement le nom du fichier, sans extension.

2. A defaut, le nom du fichier lui-meme, qui suit une regle stable :

       BCH_PLATINE_82_0889_0877-0000-CL
       └┬┘ └──────┬──────┘ └────┬────┘
      prefixe   machine       numero de liste

Le prefixe du fichier n'est pas toujours le nom du groupe : le groupe BSH
ecrit ses fichiers `BCH_…`. La table ci-dessous corrige ce seul cas connu.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SHEET = "Repertoire_LFT"
FALLBACK_SHEETS = ("Repertoire_LFT", "Repertoire", "Feuil1")

# Prefixe de nom de fichier -> groupe reel, quand les deux different.
PREFIX_TO_GROUP = {"BCH": "BSH"}

RE_LIST = re.compile(r"_(\d{3,5}-\d{3,5}-[A-Z0-9_]{1,4})$", re.I)


@dataclass(frozen=True)
class Entry:
    groupe: str
    machine: str
    description: str = ""
    source: str = "répertoire"

    @property
    def known(self) -> bool:
        return self.source == "répertoire"


def _norm(v: object) -> str:
    return str(v or "").strip()


def parse_filename(stem: str) -> Entry:
    """Deduit (groupe, machine) du seul nom de fichier. Toujours un resultat."""
    stem = _norm(stem)
    prefix = stem.split("_", 1)[0].upper() if "_" in stem else stem.upper()
    groupe = PREFIX_TO_GROUP.get(prefix, prefix)
    rest = stem[len(prefix) + 1:] if "_" in stem else ""
    m = RE_LIST.search(rest)
    machine = rest[: m.start()] if m else rest
    return Entry(groupe or "SANS_GROUPE", machine or "SANS_MACHINE", "",
                 source="nom de fichier")


def list_number_from(stem: str) -> str:
    """'BCH_PLATINE_82_0889_0877-0000-CL' -> '0877-0000-CL'."""
    m = RE_LIST.search(_norm(stem))
    return m.group(1) if m else ""


class Registry:
    """Table LFT -> (groupe, machine, description)."""

    def __init__(self, entries: dict[str, Entry] | None = None,
                 path: str | None = None):
        self.entries = entries or {}
        self.path = path

    def __len__(self) -> int:
        return len(self.entries)

    @classmethod
    def load(cls, path: str | Path | None) -> "Registry":
        """Lit le repertoire consolide. Sans fichier, retourne un registre vide
        qui se rabattra sur les noms de fichiers."""
        if not path:
            return cls()
        p = Path(path).expanduser()
        if not p.exists():
            raise FileNotFoundError(f"répertoire machines introuvable : {p}")

        from openpyxl import load_workbook

        wb = load_workbook(p, data_only=True, read_only=True)
        ws = None
        for name in FALLBACK_SHEETS:
            if name in wb.sheetnames:
                ws = wb[name]
                break
        if ws is None:
            ws = wb.worksheets[0]

        entries: dict[str, Entry] = {}
        header_seen = False
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 3:
                continue
            groupe, machine, lft = (_norm(row[0]), _norm(row[1]), _norm(row[2]))
            desc = _norm(row[3]) if len(row) > 3 else ""
            if not lft:
                continue
            if not header_seen and groupe.lower() == "groupe":
                header_seen = True
                continue
            if lft.lower() in ("lft", "code lft"):
                continue
            entries.setdefault(lft, Entry(groupe or "SANS_GROUPE",
                                          machine or "SANS_MACHINE", desc))
        wb.close()
        return cls(entries, str(p))

    def resolve(self, source: str | Path) -> Entry:
        """Rattache un fichier LFT. Ne leve jamais : un fichier inconnu tombe
        dans son groupe deduit du nom, jamais dans un fourre-tout."""
        stem = Path(source).stem
        hit = self.entries.get(stem)
        if hit:
            return hit
        # tolerance : certains exports ajoutent un suffixe de version
        for key, entry in self.entries.items():
            if stem.startswith(key):
                return entry
        return parse_filename(stem)


SAFE = re.compile(r"[^A-Za-z0-9._@+-]+")


def safe_name(value: object, default: str = "SANS_NOM", maxlen: int = 80) -> str:
    """Nom de dossier ou de fichier utilisable sous Windows comme sous Linux."""
    s = SAFE.sub("_", _norm(value)).strip("._")
    s = s[:maxlen].rstrip("._")
    return s or default


__all__ = ["Registry", "Entry", "parse_filename", "list_number_from", "safe_name",
           "PREFIX_TO_GROUP"]
